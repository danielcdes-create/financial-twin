import csv
import io
from pathlib import Path
from typing import Any
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from zipfile import BadZipFile
from xml.etree.ElementTree import ParseError

class ImportFinancialProvider:
    """Read statement rows without writing to MongoDB or guessing column meaning."""
    def read_rows(self,filename:str,data:bytes)->list[dict[str,Any]]:
        ext=Path(filename).suffix.casefold()
        if ext==".csv":
            try: text=data.decode("utf-8-sig",errors="strict")
            except UnicodeDecodeError as exc: raise ValueError("CSV must use UTF-8 encoding") from exc
            try: dialect=csv.Sniffer().sniff(text[:4096],delimiters=";,\t")
            except csv.Error: dialect=csv.excel
            rows=csv.DictReader(io.StringIO(text),dialect=dialect)
            if not rows.fieldnames: raise ValueError("CSV header row is missing")
            return list(rows)
        if ext==".xlsx":
            book=None
            try:
                book=load_workbook(io.BytesIO(data),read_only=True,data_only=True)
                sheet=book.active
                if sheet is None: raise ValueError('XLSX sem planilha legível.')
                iterator=sheet.iter_rows(values_only=True)
                headers=[str(x or "").strip() for x in next(iterator,())]
                if not any(headers): raise ValueError("XLSX header row is missing")
                return [dict(zip(headers,row)) for row in iterator]
            except (BadZipFile,InvalidFileException,ParseError,KeyError,EOFError) as exc:
                raise ValueError('O arquivo XLSX está inválido ou corrompido. Escolha uma cópia válida.') from exc
            finally:
                if book is not None: book.close()
        raise ValueError("Only CSV and XLSX imports are supported")
